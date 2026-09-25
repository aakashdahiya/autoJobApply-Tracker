"""Render the same selection to DOCX.

Some Canadian recruitment agencies and older Workday/Taleo instances still ask
for Word, and a few parse it more reliably than PDF. Same ATS rules as the PDF
path: single column, no tables, no text boxes, no header or footer, contact
details in the body.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from resume.schema import Profile
from resume.select import Selection
from resume.typst_render import build_payload

FONT = "Arial"  # metric-compatible with Liberation Sans, present everywhere


def _style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)


def _para(doc: Document, text: str = "", *, bold: bool = False, size: int = 10,
          space_before: int = 0, space_after: int = 2) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT


def _section(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = FONT
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def render_docx(profile: Profile, selection: Selection, out_path: str | Path) -> Path:
    data = build_payload(profile, selection)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    for s in doc.sections:  # generous but standard margins; no header or footer
        s.top_margin = s.bottom_margin = Pt(43)
        s.left_margin = s.right_margin = Pt(50)
    _style(doc)

    _para(doc, data["identity"]["name"], bold=True, size=17, space_after=4)
    _para(doc, "  |  ".join(data["contact"]), size=9.5)
    if data["work_auth_line"]:
        _para(doc, data["work_auth_line"], size=9.5)

    if data["summary"]:
        _section(doc, "Summary")
        _para(doc, data["summary"])

    if data["skills"]:
        _section(doc, "Skills")
        for group in data["skills"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            label = p.add_run(f"{group['category']}: ")
            label.bold = True
            label.font.name = FONT
            label.font.size = Pt(10)
            body = p.add_run(", ".join(group["items"]))
            body.font.name = FONT
            body.font.size = Pt(10)

    if data["experience"]:
        _section(doc, "Experience")
        for role in data["experience"]:
            header = f"{role['title']}, {role['company']} — {role['location']}  |  {role['dates']}"
            _para(doc, header, bold=True, space_before=6, space_after=2)
            for bullet in role["bullets"]:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(bullet)
                run.font.name = FONT
                run.font.size = Pt(10)

    if data["education"]:
        _section(doc, "Education")
        for item in data["education"]:
            header = (
                f"{item['credential']}, {item['institution']} — {item['location']}"
                f"  |  {item['dates']}"
            )
            _para(doc, header, bold=True, space_before=4, space_after=1)
            if item["note"]:
                _para(doc, item["note"], size=9.5)

    doc.save(str(out))
    return out
