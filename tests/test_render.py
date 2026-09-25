"""Rendering, and the extraction checks that make the PDF trustworthy."""

from __future__ import annotations

from docx import Document
from pdfminer.high_level import extract_text

from resume.docx_render import render_docx
from resume.render import main
from resume.schema import Shape
from resume.select import select
from resume.typst_render import render_pdf
from resume.verify import TRUE_ORDER, normalise, reading_order_risks, verify_pdf


def test_every_shape_renders_and_survives_extraction(profile, tmp_path):
    """The regression test for the whole ATS-safety claim.

    It once failed for real: right-aligning the dates made the extractor emit
    the date in the middle of the first bullet. Any layout change that reorders
    the text layer fails here rather than in a recruiter's inbox.
    """
    for shape in Shape:
        for date_style in ("tab", "inline"):
            selection = select(profile, shape)
            pdf = render_pdf(
                profile, selection, tmp_path / f"{shape.value}-{date_style}.pdf",
                date_style=date_style,
            )
            assert pdf.exists() and pdf.stat().st_size > 0
            assert verify_pdf(pdf, profile, selection) == []


def test_inline_dates_remove_the_naive_parser_risk(profile, tmp_path):
    """The escape hatch has to actually work.

    With dates at the right margin the PDF is correct but a parser that groups
    by proximity can read one into the following bullet. Inline dates leave
    nothing to misgroup, and this asserts it rather than assuming it.
    """
    selection = select(profile, Shape.fullstack)
    inline = render_pdf(profile, selection, tmp_path / "inline.pdf", date_style="inline")
    assert reading_order_risks(inline, selection) == []


def test_pdf_carries_contact_details(profile, tmp_path):
    selection = select(profile, Shape.ai_engineer)
    pdf = render_pdf(profile, selection, tmp_path / "r.pdf")

    text = normalise(extract_text(str(pdf), laparams=TRUE_ORDER))
    assert normalise(profile.identity.name) in text
    assert normalise(profile.identity.email) in text
    assert normalise(selection.headline) in text


def test_no_work_auth_line_by_default(profile, tmp_path):
    """Work authorisation is answered on the form, not advertised on the resume."""
    selection = select(profile, Shape.ai_engineer)
    pdf = render_pdf(profile, selection, tmp_path / "r.pdf")
    text = normalise(extract_text(str(pdf), laparams=TRUE_ORDER))
    assert normalise("Authorised to work in Canada") not in text


def test_work_auth_line_can_be_opted_into(profile, tmp_path):
    selection = select(profile, Shape.ai_engineer)
    pdf = render_pdf(profile, selection, tmp_path / "r.pdf", work_auth_line=True)
    text = normalise(extract_text(str(pdf), laparams=TRUE_ORDER))
    assert normalise("Authorised to work in Canada") in text


def test_pdf_has_no_repeated_header_artifacts(profile, tmp_path):
    """Contact details belong in the body; a repeated page header breaks parsers."""
    selection = select(profile, Shape.fullstack)
    pdf = render_pdf(profile, selection, tmp_path / "r.pdf")
    text = normalise(extract_text(str(pdf), laparams=TRUE_ORDER))
    assert text.count(normalise(profile.identity.email)) == 1


def test_docx_matches_the_pdf_content(profile, tmp_path):
    selection = select(profile, Shape.fullstack)
    path = render_docx(profile, selection, tmp_path / "r.docx")

    doc = Document(str(path))
    text = normalise(" ".join(p.text for p in doc.paragraphs))
    for bullet in selection.bullets:
        assert normalise(bullet.text) in text
    assert not doc.tables, "tables break ATS parsing"


def test_docx_preserves_inline_bold(profile, tmp_path):
    """Bolding must not drift between the PDF and DOCX paths."""
    selection = select(profile, Shape.fullstack)
    path = render_docx(profile, selection, tmp_path / "r.docx")

    doc = Document(str(path))
    bolded = {
        r.text
        for p in doc.paragraphs
        for r in p.runs
        if r.bold and p.style is not None and p.style.name == "List Bullet"
    }
    assert "35%" in bolded


def test_cli_renders_and_reports(tmp_path, capsys):
    out = tmp_path / "out.pdf"
    code = main(
        [
            "--profile", "profile.example.yaml",
            "--shape", "fullstack",
            "--out", str(out),
            "--docx", str(tmp_path / "out.docx"),
        ]
    )
    assert code == 0
    assert out.exists()
    assert "traceability ok, PDF text layer verified" in capsys.readouterr().out


def test_cli_depth_only_needs_no_shape(capsys):
    assert main(["--profile", "profile.example.yaml", "--depth-only"]) == 0
    assert "Fact bank depth by shape" in capsys.readouterr().out
